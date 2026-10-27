"""Export human review decisions to Excel workbook."""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment


def export_review_decisions(records_dir: str, output_path: str) -> None:
    """Export review decisions across all records to an Excel workbook."""
    records_path = Path(records_dir)
    records = []

    # Load all record files
    for record_file in sorted(records_path.glob("*.json")):
        try:
            with open(record_file, "r") as f:
                record = json.load(f)
                records.append(record)
        except Exception as e:
            print(f"Error loading {record_file}: {e}")

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Review_Decisions sheet
    ws_decisions = wb.create_sheet("Review_Decisions")
    _populate_decisions_sheet(ws_decisions, records)

    # Summary sheet
    ws_summary = wb.create_sheet("Summary")
    _populate_summary_sheet(ws_summary, records)

    wb.save(output_path)
    print(f"Exported review decisions to {output_path}")


def _populate_decisions_sheet(ws, records):
    """Populate the Review_Decisions sheet with per-item decisions."""
    # Header row
    headers = [
        "Record ID",
        "Case Number",
        "Client Name",
        "Employee Name",
        "Item ID",
        "Item Type",
        "Code",
        "Description",
        "Quantity (Value)",
        "Quantity (Unit)",
        "Rate",
        "Amount",
        "Status",
        "Confidence",
        "Decision",
        "Corrected Value",
        "Note",
        "Decided By",
        "Decided At",
    ]
    ws.append(headers)

    # Style header
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data rows
    row_num = 2
    for record in records:
        record_id = record.get("record_id", "")
        source = record.get("source", {})
        case_number = source.get("case_number", "")

        # Get client name
        document = record.get("document", {})
        client = document.get("client", {})
        if isinstance(client, dict):
            client_name = client.get("name", {}).get("value", "Unknown")
        else:
            client_name = "Unknown"

        # Get employees for lookup
        employees = {e["employee_ref"]: e for e in record.get("employees", [])}

        # Get human_review decisions
        human_review = record.get("human_review", {})
        field_decisions = human_review.get("field_decisions", [])

        # If no decisions, still list the items with no decision
        pay_items = record.get("pay_items", [])
        decided_item_ids = {d["item_id"] for d in field_decisions}

        for item in pay_items:
            item_id = item.get("item_id", "")
            emp_ref = item.get("employee_ref", "")
            emp = employees.get(emp_ref, {})

            if isinstance(emp.get("name"), dict):
                emp_name = emp.get("name", {}).get("full", "Unknown")
            else:
                emp_name = emp.get("name", "Unknown") if emp else "Unknown"

            item_type = item.get("item_type", "")
            code = item.get("code", "")
            description = item.get("code_description", "")
            quantity = item.get("quantity", {})
            quantity_val = quantity.get("value", "") if isinstance(quantity, dict) else ""
            quantity_unit = quantity.get("unit", "") if isinstance(quantity, dict) else ""
            rate = item.get("rate", "")
            amount = item.get("amount", "")
            status = item.get("status", "")
            confidence = item.get("confidence", "")

            # Find decision for this item
            decision = None
            for d in field_decisions:
                if d.get("item_id") == item_id:
                    decision = d
                    break

            decision_str = decision.get("decision", "") if decision else ""
            corrected_val = decision.get("corrected_value", "") if decision else ""
            note = decision.get("note", "") if decision else ""
            decided_by = decision.get("decided_by", "") if decision else ""
            decided_at = decision.get("decided_at", "") if decision else ""

            ws.append(
                [
                    record_id,
                    case_number,
                    client_name,
                    emp_name,
                    item_id,
                    item_type,
                    code,
                    description,
                    quantity_val,
                    quantity_unit,
                    rate,
                    amount,
                    status,
                    confidence,
                    decision_str,
                    json.dumps(corrected_val) if corrected_val else "",
                    note,
                    decided_by,
                    decided_at,
                ]
            )
            row_num += 1

    # Auto-fit columns
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)


def _populate_summary_sheet(ws, records):
    """Populate the Summary sheet with aggregate statistics."""
    total_records = len(records)
    records_complete = sum(
        1 for r in records if r.get("human_review", {}).get("status") == "complete"
    )
    records_in_progress = sum(
        1 for r in records if r.get("human_review", {}).get("status") == "in_progress"
    )
    records_not_started = total_records - records_complete - records_in_progress

    total_pay_items = sum(len(r.get("pay_items", [])) for r in records)
    total_decisions = sum(
        len(r.get("human_review", {}).get("field_decisions", [])) for r in records
    )
    accepted = sum(
        sum(1 for d in r.get("human_review", {}).get("field_decisions", [])
            if d.get("decision") == "accept")
        for r in records
    )
    rejected = sum(
        sum(1 for d in r.get("human_review", {}).get("field_decisions", [])
            if d.get("decision") == "reject")
        for r in records
    )
    edited = sum(
        sum(1 for d in r.get("human_review", {}).get("field_decisions", [])
            if d.get("decision") == "edit")
        for r in records
    )

    data = [
        ["Metric", "Value"],
        ["Total Records", total_records],
        ["Records Complete", records_complete],
        ["Records In Progress", records_in_progress],
        ["Records Not Started", records_not_started],
        ["", ""],
        ["Total Pay Items", total_pay_items],
        ["Total Decisions Made", total_decisions],
        ["Accepted", accepted],
        ["Rejected", rejected],
        ["Edited", edited],
        ["", ""],
        ["Export Date", datetime.now().isoformat()],
    ]

    for row in data:
        ws.append(row)

    # Style header
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    # Highlight metric names
    for row_num in [2, 3, 4, 5, 7, 8, 9, 10, 11, 13]:
        ws[f"A{row_num}"].font = Font(bold=True)

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 15
