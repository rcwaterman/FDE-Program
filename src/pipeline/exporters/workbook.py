"""Flatten payroll_extractions.jsonl (v2) into a review workbook.

Called by the export stage (stages/export.py); skip it with `run --no-workbook`.
One tab per v2 array; every row carries Record ID. Documents tab has formula columns that
count/sum Pay_Items by record.
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR = PatternFill("solid", fgColor="1F3864")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS_FILL = {"excluded": "F8CBAD", "superseded": "F8CBAD", "struck_through": "EDEDED",
               "added_from_email": "DDEBF7", "printed_default": "FFF2CC"}
RESULT_FILL = {"match": "E2EFDA", "corroborated": "E2EFDA", "filled_from_salesforce": "DDEBF7",
               "mismatch": "F8CBAD", "superseded": "F8CBAD", "not_on_document": "FFF2CC"}


def v(x, k="value"):
    return (x or {}).get(k) if isinstance(x, dict) else x


def sheet(wb, name, cols, rows, widths=None, fills=None):
    ws = wb.create_sheet(name)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR; cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, 14)
    for r, row in enumerate(rows, 2):
        for i, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.font = Font(name=FONT, size=10); cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(val, str) and len(val) > 40)
            if isinstance(val, float):
                cell.number_format = "#,##0.00##"
        if fills:
            col, mp = fills
            f = mp.get(row[col])
            if f:
                ws.cell(row=r, column=col + 1).fill = PatternFill("solid", fgColor=f)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(rows)) + 1}"
    return ws


def main(src, out):
    recs = [json.loads(l) for l in open(src) if l.strip()]
    wb = Workbook(); wb.remove(wb.active)

    docs = []
    for r in recs:
        d = r["document"]; pp = d.get("pay_period") or {}; cl = d.get("client") or {}
        pc = next((c for c in d.get("contacts", []) if c["role"] == "payroll_contact"), {})
        sv = r.get("salesforce_verification") or {}
        docs.append([r["record_id"], r["source"]["case_number"], r["source"]["page_count"], d["doc_type"], d.get("form_variant"),
                     v(cl.get("name")), v(cl.get("number")), v(pp.get("start")), v(pp.get("end")), v(pp.get("check_date")),
                     v(pp.get("frequency")), pc.get("name"), pc.get("phone"),
                     (sv.get("account") or {}).get("name"), (sv.get("account") or {}).get("employer_unique_id"),
                     (sv.get("account") or {}).get("payroll_frequency"), len(r["employees"]),
                     (r.get("extraction") or {}).get("batch"), None, None, None, None, None])
    cols = ["Record ID", "Case", "Pages", "Doc Type", "Form Variant", "Client Name", "Client #", "Period Start", "Period End",
            "Check Date", "Freq", "Payroll Contact", "Contact Phone", "SF Account", "SF Client #", "SF Frequency", "Employees",
            "Batch", "Pay Items", "Payable Hours", "Payable Amount", "Totals Not Reconciled", "Review Flags"]
    ws = sheet(wb, "Documents", cols, docs, {"Record ID": 30, "Form Variant": 30, "Client Name": 28, "SF Account": 28})
    for i in range(2, len(docs) + 2):
        ws[f"S{i}"] = f"=COUNTIF(Pay_Items!$A:$A,$A{i})"
        ws[f"T{i}"] = f'=SUMIFS(Pay_Items!$I:$I,Pay_Items!$A:$A,$A{i},Pay_Items!$J:$J,"hours",Pay_Items!$N:$N,"<>excluded",Pay_Items!$N:$N,"<>superseded",Pay_Items!$N:$N,"<>struck_through")'
        ws[f"U{i}"] = f'=SUMIFS(Pay_Items!$L:$L,Pay_Items!$A:$A,$A{i},Pay_Items!$N:$N,"<>excluded",Pay_Items!$N:$N,"<>superseded",Pay_Items!$N:$N,"<>struck_through")'
        ws[f"V{i}"] = f'=COUNTIFS(Totals!$A:$A,$A{i},Totals!$J:$J,FALSE)'
        ws[f"W{i}"] = f"=COUNTIF(Review_Flags!$A:$A,$A{i})"
        for c in "STUVW":
            ws[f"{c}{i}"].font = Font(name=FONT, size=10); ws[f"{c}{i}"].border = BORDER

    emp = [[r["record_id"], e.get("page"), e["employee_ref"], e["name"]["full"], e.get("employee_number"), e.get("clock_number"),
            e.get("pay_type"), e.get("status"), json.dumps(e.get("attributes") or {}, ensure_ascii=False)]
           for r in recs for e in r["employees"]]
    sheet(wb, "Employees", ["Record ID", "Page", "Ref", "Name", "Emp #", "Clock #", "Pay Type", "Status", "Attributes (JSON)"],
          emp, {"Record ID": 30, "Name": 28, "Attributes (JSON)": 90})

    pi = []
    for r in recs:
        for p in r["pay_items"]:
            q = p.get("quantity") or {}; pv = p["provenance"]
            pi.append([r["record_id"], pv.get("page"), p["item_id"], p.get("employee_ref"), p.get("employee_name"), p["item_type"],
                       p.get("code"), p.get("code_description"), q.get("value"), q.get("unit"), p.get("rate"),
                       p.get("amount"), pv.get("source"), p["status"], p.get("confidence"), ", ".join(p.get("review_reasons") or []),
                       pv.get("raw_text"), pv.get("email_message_id"), p.get("comment")])
    sheet(wb, "Pay_Items", ["Record ID", "Page", "Item", "Emp Ref", "Employee", "Type", "Code", "Code Desc", "Qty", "Unit",
                            "Rate", "Amount", "Source", "Status", "Confidence", "Review Reasons", "Raw Text", "Email Msg", "Comment"],
          pi, {"Record ID": 30, "Employee": 26, "Code Desc": 18, "Comment": 50, "Review Reasons": 20}, fills=(13, STATUS_FILL))

    ded = [[r["record_id"], e["employee_ref"], e["name"]["full"], d["code"], d.get("description"), d.get("amount"), d.get("ytd"),
            d.get("balance"), d.get("change_requested")] for r in recs for e in r["employees"] for d in e.get("recurring_deductions") or []]
    sheet(wb, "Recurring_Deductions", ["Record ID", "Emp Ref", "Employee", "Code", "Description", "Amount", "YTD", "Balance", "Change Requested"],
          ded, {"Record ID": 30, "Employee": 26})

    leg = [[r["record_id"], c["code"], c["description"], c["category"]] for r in recs for c in r.get("code_legend") or []]
    sheet(wb, "Code_Legend", ["Record ID", "Code", "Description", "Category"], leg, {"Record ID": 30, "Description": 24})

    tot = [[r["record_id"], t.get("page"), t["label"], t.get("scope"), t.get("code"), t.get("measure"), t["document_value"],
            t.get("extracted_sum"), t.get("difference"), t.get("reconciles"), t.get("source")] for r in recs for t in r.get("totals") or []]
    sheet(wb, "Totals", ["Record ID", "Page", "Label", "Scope", "Code", "Measure", "Document Value", "Extracted Sum", "Difference",
                         "Reconciles", "Source"], tot, {"Record ID": 30, "Label": 36})

    tp = [[r["record_id"], t.get("page"), t.get("employee_name"), t.get("date"),
           "; ".join(f"{s.get('in')}-{s.get('out')}" for s in t.get("segments") or []), t.get("hours"), t.get("minutes"), t.get("source")]
          for r in recs for t in r.get("time_punches") or []]
    sheet(wb, "Time_Punches", ["Record ID", "Page", "Employee", "Date", "Segments", "Hours", "Minutes", "Source"], tp, {"Record ID": 30, "Segments": 30})

    dd = [[r["record_id"], x.get("page"), x.get("employee_name"), x.get("ssn"), x.get("action"), a.get("seq"), a.get("account_type"),
           a.get("routing_number"), a.get("account_number"), a.get("deposit_type"), a.get("deposit_value"), x.get("signed"), x.get("comment")]
          for r in recs for x in r.get("direct_deposit") or [] for a in x.get("accounts") or [{}]]
    sheet(wb, "Direct_Deposit", ["Record ID", "Page", "Employee", "SSN", "Action", "Seq", "Type", "Routing", "Account", "Deposit Type",
                                 "Value", "Signed", "Comment"], dd, {"Record ID": 30, "Comment": 50})

    ins = [[r["record_id"], i.get("page"), i.get("kind"), i.get("employee_ref"), i["text"]] for r in recs for i in r["document"].get("instructions") or []]
    sheet(wb, "Instructions", ["Record ID", "Page", "Kind", "Emp Ref", "Text"], ins, {"Record ID": 30, "Text": 90})

    chk = [[r["record_id"], c["field"], c.get("extracted"), c.get("salesforce"), c.get("sf_source"), c["result"], c.get("note")]
           for r in recs for c in (r.get("salesforce_verification") or {}).get("checks", [])]
    sheet(wb, "SF_Checks", ["Record ID", "Field", "Extracted", "Salesforce", "SF Source", "Result", "Note"], chk,
          {"Record ID": 30, "Field": 22, "Extracted": 34, "Salesforce": 40, "Note": 60}, fills=(5, RESULT_FILL))

    am = [[r["record_id"], a.get("message_date"), a.get("from"), a["action"], a["instruction"], ", ".join(a.get("affects") or []),
           a.get("email_message_id")] for r in recs for a in (r.get("salesforce_verification") or {}).get("email_amendments", [])]
    sheet(wb, "Email_Amendments", ["Record ID", "Date", "From", "Action", "Instruction", "Affects", "Email Msg"], am,
          {"Record ID": 30, "From": 28, "Instruction": 80})

    fl = [[r["record_id"], f.get("code"), f["severity"], f["message"], ", ".join(f.get("refs") or [])] for r in recs for f in r["review_flags"]]
    sheet(wb, "Review_Flags", ["Record ID", "Code", "Severity", "Message", "Refs"], fl, {"Record ID": 30, "Message": 110})

    pg = [[r["record_id"], p["page"], p.get("orientation"), (p.get("printed_text") or "")[:32000], (p.get("handwritten_text") or "")[:32000]]
          for r in recs for p in r.get("pages") or []]
    ws = sheet(wb, "Pages", ["Record ID", "Page", "Orientation", "Printed Text", "Handwritten Text"], pg,
               {"Record ID": 30, "Printed Text": 90, "Handwritten Text": 70})

    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 34; sm.column_dimensions["B"].width = 14
    rows = [("Payroll Attachment Extractions (schema v2)", None), ("Source", src.split("/")[-1]), (None, None),
            ("Records", "=COUNTA(Documents!A:A)-1"), ("Pay items", "=COUNTA(Pay_Items!A:A)-1"),
            ("Pay items excluded/superseded by email", '=COUNTIF(Pay_Items!N:N,"excluded")+COUNTIF(Pay_Items!N:N,"superseded")'),
            ("Pay items added from email", '=COUNTIF(Pay_Items!N:N,"added_from_email")'),
            ("Totals reconciled", "=COUNTIF(Totals!J:J,TRUE)"), ("Totals not reconciled", "=COUNTIF(Totals!J:J,FALSE)"),
            ("SF checks", "=COUNTA(SF_Checks!A:A)-1"), ("SF match/corroborated", '=COUNTIF(SF_Checks!F:F,"match")+COUNTIF(SF_Checks!F:F,"corroborated")'),
            ("SF filled_from_salesforce", '=COUNTIF(SF_Checks!F:F,"filled_from_salesforce")'),
            ("SF mismatch", '=COUNTIF(SF_Checks!F:F,"mismatch")'), ("SF superseded", '=COUNTIF(SF_Checks!F:F,"superseded")'),
            ("SF not_on_document", '=COUNTIF(SF_Checks!F:F,"not_on_document")'), ("Email amendments", "=COUNTA(Email_Amendments!A:A)-1"),
            ("Review flags", "=COUNTA(Review_Flags!A:A)-1"), (None, None),
            ("Pay item status colours", "yellow = printed_default, grey = struck_through, red = excluded/superseded, blue = added_from_email"),
            ("Sensitive data", "Direct_Deposit and some Pages text contain SSNs / bank numbers as written.")]
    for i, (a, b) in enumerate(rows, 1):
        sm.cell(row=i, column=1, value=a).font = Font(name=FONT, bold=(i == 1), size=14 if i == 1 else 10)
        c = sm.cell(row=i, column=2, value=b); c.font = Font(name=FONT, size=10)
    wb.save(out)
    print(f"{len(recs)} records -> {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
